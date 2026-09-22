"""Convert docs/frontend-customization-guide.md into a Word .docx file.

Best-effort markdown -> docx converter tailored for the customization guide:
- Handles ATX headings (# through ####), fenced code blocks, blockquotes,
  bullet lists, numbered lists, and markdown tables.
- Preserves inline `code`, **bold**, and *italic* runs.
- Rewrites relative markdown links so they show as readable text (the target
  path shown after the label if the label differs).

Usage:
    python scripts/md_to_docx.py docs/frontend-customization-guide.md \\
        docs/frontend-customization-guide.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INLINE_CODE = re.compile(r"`([^`]+)`")
BOLD = re.compile(r"\*\*([^*]+)\*\*")
ITALIC = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _add_runs(paragraph, text: str) -> None:
    """Add text to `paragraph`, applying simple inline markdown formatting."""

    tokens: list[tuple[str, str]] = []

    def _tokenize(chunk: str, kind: str) -> None:
        if not chunk:
            return
        # Resolve links first so labels can then be inline-formatted.
        pieces: list[tuple[str, str]] = []
        pos = 0
        for match in LINK.finditer(chunk):
            if match.start() > pos:
                pieces.append((kind, chunk[pos : match.start()]))
            label, target = match.group(1), match.group(2)
            if label == target or label.strip() == target.strip():
                pieces.append((kind, label))
            else:
                pieces.append((kind, f"{label} ({target})"))
            pos = match.end()
        if pos < len(chunk):
            pieces.append((kind, chunk[pos:]))

        for piece_kind, piece_text in pieces:
            # Split on inline code first.
            last = 0
            for m in INLINE_CODE.finditer(piece_text):
                if m.start() > last:
                    _apply_bold_italic(piece_text[last : m.start()], piece_kind, tokens)
                tokens.append(("code", m.group(1)))
                last = m.end()
            if last < len(piece_text):
                _apply_bold_italic(piece_text[last:], piece_kind, tokens)

    _tokenize(text, "text")

    for kind, value in tokens:
        run = paragraph.add_run(value)
        if kind == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        if kind in ("bold", "bold-italic"):
            run.bold = True
        if kind in ("italic", "bold-italic"):
            run.italic = True


def _apply_bold_italic(text: str, base_kind: str, tokens: list[tuple[str, str]]) -> None:
    """Split `text` on **bold** and *italic* markers and append tokens."""

    remaining = text
    while remaining:
        b = BOLD.search(remaining)
        i = ITALIC.search(remaining)
        # pick the earliest match
        candidates = [(m.start(), kind, m) for kind, m in (("bold", b), ("italic", i)) if m]
        if not candidates:
            tokens.append((base_kind, remaining))
            return
        candidates.sort(key=lambda x: x[0])
        start, kind, match = candidates[0]
        if start > 0:
            tokens.append((base_kind, remaining[:start]))
        combined = kind
        if base_kind == "italic" and kind == "bold":
            combined = "bold-italic"
        if base_kind == "bold" and kind == "italic":
            combined = "bold-italic"
        tokens.append((combined, match.group(1)))
        remaining = remaining[match.end() :]


def _iter_blocks(lines: list[str]) -> Iterable[tuple[str, list[str]]]:
    """Group markdown lines into (block_type, lines) tuples."""

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            code: list[str] = []
            while j < len(lines) and not lines[j].startswith("```"):
                code.append(lines[j])
                j += 1
            yield ("code", code)
            i = j + 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("#"):
            yield ("heading", [line])
            i += 1
            continue
        if line.startswith("> "):
            block = [line[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("> "):
                block.append(lines[i][2:])
                i += 1
            yield ("quote", block)
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|\s*$", lines[i + 1]):
            block = [line]
            i += 1
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            yield ("table", block)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            block = [line]
            i += 1
            while i < len(lines) and (re.match(r"^\s*[-*]\s+", lines[i]) or (lines[i].startswith("  ") and lines[i].strip())):
                block.append(lines[i])
                i += 1
            yield ("ul", block)
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            block = [line]
            i += 1
            while i < len(lines) and (re.match(r"^\s*\d+\.\s+", lines[i]) or (lines[i].startswith("  ") and lines[i].strip())):
                block.append(lines[i])
                i += 1
            yield ("ol", block)
            continue
        # paragraph
        block = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not _is_block_start(lines[i]):
            block.append(lines[i])
            i += 1
        yield ("p", block)


def _is_block_start(line: str) -> bool:
    if line.startswith("#") or line.startswith("```") or line.startswith("> "):
        return True
    if line.startswith("|"):
        return True
    if re.match(r"^\s*[-*]\s+", line):
        return True
    if re.match(r"^\s*\d+\.\s+", line):
        return True
    return False


def _add_code_block(doc: Document, code_lines: list[str]) -> None:
    p = doc.add_paragraph()
    run = p.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def _add_table(doc: Document, rows: list[str]) -> None:
    def split_row(r: str) -> list[str]:
        return [c.strip() for c in r.strip().strip("|").split("|")]

    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for idx, cell_text in enumerate(header):
        cell_para = hdr_cells[idx].paragraphs[0]
        run = cell_para.add_run(cell_text)
        run.bold = True
    for row_idx, row_values in enumerate(body, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, cell_value in enumerate(row_values):
            if col_idx >= len(row_cells):
                continue
            _add_runs(row_cells[col_idx].paragraphs[0], cell_value)
    doc.add_paragraph()


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Base font tweaks.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for block_type, block_lines in _iter_blocks(lines):
        if block_type == "heading":
            raw = block_lines[0]
            level = len(raw) - len(raw.lstrip("#"))
            text = raw[level:].strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(22)
            else:
                doc.add_heading(text, level=min(level - 1, 4))
        elif block_type == "code":
            _add_code_block(doc, block_lines)
        elif block_type == "table":
            _add_table(doc, block_lines)
        elif block_type == "ul":
            for item in block_lines:
                text = re.sub(r"^\s*[-*]\s+", "", item)
                p = doc.add_paragraph(style="List Bullet")
                _add_runs(p, text)
        elif block_type == "ol":
            for item in block_lines:
                text = re.sub(r"^\s*\d+\.\s+", "", item)
                p = doc.add_paragraph(style="List Number")
                _add_runs(p, text)
        elif block_type == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_runs(p, " ".join(block_lines).strip())
            for run in p.runs:
                run.italic = True
        else:  # paragraph
            joined = " ".join(l.strip() for l in block_lines)
            if set(joined.strip()) == {"-"} and len(joined.strip()) >= 3:
                # horizontal rule
                doc.add_paragraph("_" * 40)
                continue
            p = doc.add_paragraph()
            _add_runs(p, joined)

    doc.save(str(docx_path))


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("usage: md_to_docx.py <input.md> <output.docx>", file=sys.stderr)
        return 2
    md_path = Path(argv[1])
    docx_path = Path(argv[2])
    if not md_path.exists():
        print(f"input not found: {md_path}", file=sys.stderr)
        return 1
    convert(md_path, docx_path)
    print(f"wrote {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
